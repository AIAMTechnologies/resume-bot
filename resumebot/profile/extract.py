"""Turn files into plain text."""
from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

TEXT_EXTS = {".txt", ".md", ".markdown", ".rst", ".json", ".yaml", ".yml", ".csv", ".html", ".htm"}


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def pdf_text(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def docx_text(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def file_text(path: Path, limit: int = 200_000) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        text = pdf_text(path)
    elif ext == ".docx":
        text = docx_text(path)
    elif ext in TEXT_EXTS:
        text = path.read_text(errors="ignore")
    else:
        return ""
    if ext in {".html", ".htm"}:
        from bs4 import BeautifulSoup
        text = BeautifulSoup(text, "html.parser").get_text("\n")
    return text[:limit]


def unzip(path: Path, dest: Path) -> Path:
    dest.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path) as z:
        for member in z.infolist():
            target = (dest / member.filename).resolve()
            if not str(target).startswith(str(dest.resolve())):
                continue  # zip-slip guard
            z.extract(member, dest)
    return dest
