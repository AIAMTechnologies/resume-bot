"""ATS checks: parse the rendered file back and score it like an ATS would."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

ALIASES: dict[str, list[str]] = {
    "javascript": ["js", "ecmascript"],
    "typescript": ["ts"],
    "postgresql": ["postgres", "psql"],
    "kubernetes": ["k8s"],
    "amazon web services": ["aws"],
    "google cloud platform": ["gcp", "google cloud"],
    "microsoft azure": ["azure"],
    "continuous integration": ["ci/cd", "ci", "cicd"],
    "machine learning": ["ml"],
    "artificial intelligence": ["ai"],
    "large language models": ["llm", "llms"],
    "node.js": ["node", "nodejs"],
    "react": ["react.js", "reactjs"],
    "next.js": ["nextjs", "next"],
    "rest api": ["rest", "restful", "rest apis", "restful apis"],
    "c#": ["csharp", ".net"],
    "user experience": ["ux"],
    "user interface": ["ui"],
}
REQUIRED_SECTIONS = ["experience", "education", "skills"]
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(\+?\d[\d\s().-]{8,}\d)")


def normalize(text: str) -> str:
    return " " + re.sub(r"[^a-z0-9+#./]+", " ", text.lower()) + " "


def _variants(keyword: str) -> list[str]:
    k = keyword.lower().strip()
    out = {k}
    for canon, alts in ALIASES.items():
        if k == canon or k in alts:
            out.add(canon)
            out.update(alts)
    return [normalize(v).strip() for v in out if v]


def has_keyword(norm_text: str, keyword: str) -> bool:
    return any(f" {v} " in norm_text for v in _variants(keyword))


def coverage(text: str, keywords: list[str]) -> tuple[float, list[str], list[str]]:
    if not keywords:
        return 1.0, [], []
    nt = normalize(text)
    hit = [k for k in keywords if has_keyword(nt, k)]
    miss = [k for k in keywords if k not in hit]
    return len(hit) / len(keywords), hit, miss


@dataclass
class ATSReport:
    score: int
    keyword_coverage: float
    matched: list[str] = field(default_factory=list)
    missing: list[str] = field(default_factory=list)
    issues: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not any(i.startswith("BLOCKER") for i in self.issues)


def parse_back(path: Path) -> str:
    from ..profile.extract import file_text
    return file_text(path)


def check(pdf: Path, docx: Path | None, keywords: list[str], contact: dict) -> ATSReport:
    issues: list[str] = []
    text = parse_back(pdf)
    if len(text.strip()) < 300:
        issues.append("BLOCKER: PDF text is not extractable")
    lower = text.lower()
    for section in REQUIRED_SECTIONS:
        if section not in lower:
            issues.append(f"Missing standard section heading: {section.title()}")
    if contact.get("email") and contact["email"].lower() not in lower:
        issues.append("BLOCKER: email not found in parsed text")
    elif not EMAIL_RE.search(text):
        issues.append("No email address found")
    if not PHONE_RE.search(text):
        issues.append("No phone number found")

    from pypdf import PdfReader
    pages = len(PdfReader(str(pdf)).pages)
    if pages > 2:
        issues.append(f"Resume is {pages} pages (aim for 1–2)")

    if docx and docx.exists():
        import docx as docxlib
        d = docxlib.Document(str(docx))
        if d.tables:
            issues.append("DOCX contains tables (some ATS scramble them)")
        if d.inline_shapes:
            issues.append("DOCX contains images")
        if any(s.header.paragraphs and any(p.text.strip() for p in s.header.paragraphs) for s in d.sections):
            issues.append("Content in DOCX header (often skipped by ATS)")

    cov, hit, miss = coverage(text, keywords)
    penalty = sum(20 if i.startswith("BLOCKER") else 5 for i in issues)
    score = max(0, min(100, round(cov * 70 + 30 - penalty)))
    return ATSReport(score=score, keyword_coverage=round(cov, 3), matched=hit, missing=miss, issues=issues)
